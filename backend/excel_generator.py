# backend/excel_generator.py
"""
Excel Report Generation Module:
- High-resolution Logo + 'إدارة نظم المعلومات والتحول الرقمي' Header Banner
- Master Flat Table for HR Payroll analysis & Pivot Tables
- Doctor Summary Sheet with KPIs & Totals
- Highlighting for Deficit/Late (Peach #FAD9B5) and Overtime (Mint #D4EDDA)
- Footer attribution: 'Created with Eng. Mohamed Marei - إدارة نظم المعلومات والتحول الرقمي'
"""

import os
from typing import Dict, Any, List
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage


# Styling definitions
PEACH_FILL = PatternFill(start_color="FAD9B5", end_color="FAD9B5", fill_type="solid")
MINT_FILL = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
WARNING_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
HEADER_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
DARK_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
DEPT_HEADER_FILL = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

FONT_WHITE_BOLD_LARGE = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
FONT_WHITE_BOLD = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True)
FONT_NORMAL = Font(name="Calibri", size=10)

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

ALIGN_CENTER = Alignment(horizontal='center', vertical='center')
ALIGN_RIGHT = Alignment(horizontal='right', vertical='center')

LOGO_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "storage", "assets", "logo.jpeg"))
if not os.path.exists(LOGO_PATH):
    LOGO_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "frontend", "src", "logo.jpeg"))


def generate_excel_report(data: Dict[str, Any], output_path: str) -> str:
    """
    Builds the comprehensive multi-tab Excel Workbook from processed doctor attendance data.
    """
    wb = Workbook()

    doctors: List[Dict[str, Any]] = data.get('doctors', [])
    summary: Dict[str, Any] = data.get('summary', {})

    # -------------------------------------------------------------
    # SHEET 1: Master Attendance Data (Flat Table for HR Analysis)
    # -------------------------------------------------------------
    ws_master = wb.active
    ws_master.title = "السجل التفصيلي المجمع"
    ws_master.views.sheetView[0].rightToLeft = True

    # 1. Top Organization Header Banner
    ws_master.merge_cells("A1:P1")
    c1 = ws_master.cell(row=1, column=1, value="شركة الخدمات الطبية - مستشفى الكهرباء")
    c1.font = FONT_WHITE_BOLD_LARGE
    c1.fill = DARK_HEADER_FILL
    c1.alignment = ALIGN_CENTER

    ws_master.merge_cells("A2:P2")
    c2 = ws_master.cell(row=2, column=1, value="إدارة نظم المعلومات والتحول الرقمي")
    c2.font = FONT_WHITE_BOLD
    c2.fill = DEPT_HEADER_FILL
    c2.alignment = ALIGN_CENTER

    ws_master.merge_cells("A3:P3")
    c3 = ws_master.cell(row=3, column=1, value="سجل الحضور والانصراف التفصيلي للأطباء والعاملين")
    c3.font = Font(name="Calibri", size=10, bold=True, color="1B365D")
    c3.fill = HEADER_FILL
    c3.alignment = ALIGN_CENTER

    # Attach Logo image if available
    if os.path.exists(LOGO_PATH):
        try:
            xl_logo = XLImage(LOGO_PATH)
            xl_logo.width = 46
            xl_logo.height = 46
            ws_master.add_image(xl_logo, "A1")
        except Exception:
            pass

    headers = [
        "كود الموظف",
        "اسم الموظف",
        "الوظيفة",
        "الإدارة / مركز التكلفة",
        "نوع العقد",
        "التاريخ",
        "اليوم",
        "حضور",
        "انصراف",
        "حضور شفت مسائي",
        "انصراف شفت مسائي",
        "الوقت المقضي",
        "نوع الشفت",
        "وقت التأخير",
        "الوقت الإضافي",
        "ملاحظات البصمة"
    ]

    # Write Headers at Row 5
    for col_idx, h in enumerate(headers, 1):
        cell = ws_master.cell(row=5, column=col_idx, value=h)
        cell.font = FONT_WHITE_BOLD
        cell.fill = DARK_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    current_row = 6
    for doc in doctors:
        emp_code = doc['emp_code']
        emp_name = doc['emp_name']
        job_title = doc['job_title']
        department = doc['department']
        contract = doc['contract']

        for r in doc.get('records', []):
            ws_master.cell(row=current_row, column=1, value=emp_code).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=2, value=emp_name).alignment = ALIGN_RIGHT
            ws_master.cell(row=current_row, column=3, value=job_title).alignment = ALIGN_RIGHT
            ws_master.cell(row=current_row, column=4, value=department).alignment = ALIGN_RIGHT
            ws_master.cell(row=current_row, column=5, value=contract).alignment = ALIGN_RIGHT
            ws_master.cell(row=current_row, column=6, value=r.get('date', '')).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=7, value=r.get('day', '')).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=8, value=r.get('check_in_1', '')).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=9, value=r.get('check_out_1', '')).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=10, value=r.get('check_in_2', '')).alignment = ALIGN_CENTER
            ws_master.cell(row=current_row, column=11, value=r.get('check_out_2', '')).alignment = ALIGN_CENTER

            c_spent = ws_master.cell(row=current_row, column=12, value=r.get('time_spent', '0:00'))
            c_spent.alignment = ALIGN_CENTER

            c_shift = ws_master.cell(row=current_row, column=13, value=r.get('shift_type', ''))
            c_shift.alignment = ALIGN_CENTER

            c_def = ws_master.cell(row=current_row, column=14, value=r.get('deficit_time', '0:00'))
            c_def.alignment = ALIGN_CENTER
            if r.get('is_deficit'):
                c_def.fill = PEACH_FILL
                c_spent.fill = PEACH_FILL

            c_ot = ws_master.cell(row=current_row, column=15, value=r.get('overtime_time', '0:00'))
            c_ot.alignment = ALIGN_CENTER
            if r.get('is_overtime'):
                c_ot.fill = MINT_FILL

            # Column 16: Biometric Notes & Warnings
            kw = r.get('key_warning')
            note_text = ""
            if kw == 'same_key_repeat':
                note_text = f"تكرار نفس المفتاح ({r.get('in_1_key')}) للدخول والخروج"
            elif kw == 'cross_key':
                note_text = f"مفاتيح متقاطعة ({r.get('in_1_key')} ← {r.get('out_1_key')})"
            elif kw == 'reversed_keys':
                note_text = f"عكس مفاتيح الحضور والانصراف ({r.get('in_1_key')} ← {r.get('out_1_key')})"
            elif kw == 'started_with_out_key':
                note_text = f"بدء بمفتاح خروج ({r.get('in_1_key')})"
            elif kw == 'missing_checkout':
                note_text = "بصمة مفردة - لم يسجل انصراف"
            elif kw == 'missing_checkin':
                note_text = "بصمة مفردة - لم يسجل حضور"
            elif kw == 'start_of_24h':
                note_text = "بداية نوبتجية ممتدة (حضور)"
            elif kw == 'end_of_24h':
                note_text = "نهاية نوبتجية ممتدة (انصراف)"
            elif r.get('is_late_arrival'):
                note_text = "تأخير في الحضور (بعد 08:45 ص)"

            c_note = ws_master.cell(row=current_row, column=16, value=note_text)
            c_note.alignment = ALIGN_RIGHT
            if kw and kw not in ['start_of_24h', 'end_of_24h']:
                c_note.fill = WARNING_FILL
                c_shift.fill = WARNING_FILL

            for col_i in range(1, 17):
                ws_master.cell(row=current_row, column=col_i).border = THIN_BORDER
                ws_master.cell(row=current_row, column=col_i).font = FONT_NORMAL

            current_row += 1

    # Master Footer attribution
    current_row += 1
    ws_master.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=16)
    f_cell = ws_master.cell(row=current_row, column=1, value="Created with Eng. Mohamed Marei - إدارة نظم المعلومات والتحول الرقمي")
    f_cell.font = Font(name="Calibri", size=10, italic=True, color="555555")
    f_cell.alignment = ALIGN_CENTER

    # -------------------------------------------------------------
    # SHEET 2: Summary of Doctors (HR Dashboard & Totals)
    # -------------------------------------------------------------
    ws_summary = wb.create_sheet(title="ملخص الأطباء وساعات العمل")
    ws_summary.views.sheetView[0].rightToLeft = True

    # 1. Top Organization Header Banner
    ws_summary.merge_cells("A1:I1")
    s1 = ws_summary.cell(row=1, column=1, value="شركة الخدمات الطبية - مستشفى الكهرباء")
    s1.font = FONT_WHITE_BOLD_LARGE
    s1.fill = DARK_HEADER_FILL
    s1.alignment = ALIGN_CENTER

    ws_summary.merge_cells("A2:I2")
    s2 = ws_summary.cell(row=2, column=1, value="إدارة نظم المعلومات والتحول الرقمي")
    s2.font = FONT_WHITE_BOLD
    s2.fill = DEPT_HEADER_FILL
    s2.alignment = ALIGN_CENTER

    ws_summary.merge_cells("A3:I3")
    s3 = ws_summary.cell(row=3, column=1, value="تقرير ملخص الحضور، العجز، وساعات العمل الإضافية")
    s3.font = Font(name="Calibri", size=10, bold=True, color="1B365D")
    s3.fill = HEADER_FILL
    s3.alignment = ALIGN_CENTER

    if os.path.exists(LOGO_PATH):
        try:
            xl_logo2 = XLImage(LOGO_PATH)
            xl_logo2.width = 46
            xl_logo2.height = 46
            ws_summary.add_image(xl_logo2, "A1")
        except Exception:
            pass

    sum_headers = [
        "كود الموظف",
        "اسم الموظف",
        "الوظيفة",
        "الإدارة / مركز التكلفة",
        "نوع العقد",
        "أيام العمل الفعلية",
        "إجمالي الوقت المقضي",
        "إجمالي وقت التأخير / العجز",
        "إجمالي الوقت الإضافي"
    ]

    for col_idx, h in enumerate(sum_headers, 1):
        cell = ws_summary.cell(row=5, column=col_idx, value=h)
        cell.font = FONT_WHITE_BOLD
        cell.fill = DARK_HEADER_FILL
        cell.alignment = ALIGN_CENTER
        cell.border = THIN_BORDER

    s_row = 6
    for doc in doctors:
        ws_summary.cell(row=s_row, column=1, value=doc['emp_code']).alignment = ALIGN_CENTER
        ws_summary.cell(row=s_row, column=2, value=doc['emp_name']).alignment = ALIGN_RIGHT
        ws_summary.cell(row=s_row, column=3, value=doc['job_title']).alignment = ALIGN_RIGHT
        ws_summary.cell(row=s_row, column=4, value=doc['department']).alignment = ALIGN_RIGHT
        ws_summary.cell(row=s_row, column=5, value=doc['contract']).alignment = ALIGN_RIGHT
        ws_summary.cell(row=s_row, column=6, value=doc['total_days_worked']).alignment = ALIGN_CENTER
        ws_summary.cell(row=s_row, column=7, value=doc['total_time_spent']).alignment = ALIGN_CENTER
        
        c_tot_def = ws_summary.cell(row=s_row, column=8, value=doc['total_deficit'])
        c_tot_def.alignment = ALIGN_CENTER
        if doc.get('total_deficit_minutes', 0) > 0:
            c_tot_def.fill = PEACH_FILL

        c_tot_ot = ws_summary.cell(row=s_row, column=9, value=doc['total_overtime'])
        c_tot_ot.alignment = ALIGN_CENTER
        if doc.get('total_overtime_minutes', 0) > 0:
            c_tot_ot.fill = MINT_FILL

        for c in range(1, 10):
            ws_summary.cell(row=s_row, column=c).border = THIN_BORDER
            ws_summary.cell(row=s_row, column=c).font = FONT_NORMAL

        s_row += 1

    # Totals Row in Summary Sheet
    ws_summary.cell(row=s_row, column=1, value="الإجمالي العام").alignment = ALIGN_CENTER
    ws_summary.cell(row=s_row, column=1).font = FONT_BOLD
    ws_summary.cell(row=s_row, column=6, value=f"{len(doctors)} موظف").alignment = ALIGN_CENTER
    ws_summary.cell(row=s_row, column=6).font = FONT_BOLD
    ws_summary.cell(row=s_row, column=7, value=summary.get('total_hours', '0:00')).alignment = ALIGN_CENTER
    ws_summary.cell(row=s_row, column=7).font = FONT_BOLD
    ws_summary.cell(row=s_row, column=8, value=summary.get('total_deficit', '0:00')).alignment = ALIGN_CENTER
    ws_summary.cell(row=s_row, column=8).font = FONT_BOLD
    ws_summary.cell(row=s_row, column=9, value=summary.get('total_overtime', '0:00')).alignment = ALIGN_CENTER
    ws_summary.cell(row=s_row, column=9).font = FONT_BOLD

    for c in range(1, 10):
        ws_summary.cell(row=s_row, column=c).fill = HEADER_FILL
        ws_summary.cell(row=s_row, column=c).border = THIN_BORDER

    # Summary Footer attribution
    s_row += 2
    ws_summary.merge_cells(start_row=s_row, start_column=1, end_row=s_row, end_column=9)
    f_cell2 = ws_summary.cell(row=s_row, column=1, value="Created with Eng. Mohamed Marei - إدارة نظم المعلومات والتحول الرقمي")
    f_cell2.font = Font(name="Calibri", size=10, italic=True, color="555555")
    f_cell2.alignment = ALIGN_CENTER

    # Auto-adjust column widths (ignoring title row merges for width)
    for sheet in [ws_master, ws_summary]:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col[4:-1]:
                val_str = str(cell.value or '')
                max_len = max(max_len, len(val_str))
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 13)

    # Save workbook
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path
